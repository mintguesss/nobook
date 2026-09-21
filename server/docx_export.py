"""把筆記輸出成 Word 檔。

重點是**不要把 Markdown 字元原樣倒進 docx**。那樣做出來的檔案打開會看到
一堆 `##` 和 `-`，標題不是標題、清單不是清單，在 Word 裡既不能摺疊也不能
用導覽窗格跳轉，複製到別的文件也帶不走格式。

所以這裡把 Markdown 解析成結構，再對應到 Word 的原生樣式：
    # 標題    → Heading 1
    ## 小節   → Heading 2
    - 項目    → List Bullet
    > 引用    → Intense Quote
    `程式碼`  → 等寬字型的 run
粗體與行內程式碼也照樣轉成 run 層級的格式。
"""
from __future__ import annotations

import io
import re

_H_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_LI_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
# 行內格式：**粗體**、`程式碼`
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")

CJK_FONT = "Microsoft JhengHei"


def _add_runs(par, text: str):
    """把一行文字加進段落，處理 **粗體** 與 `程式碼`。"""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part)


def _set_cjk_font(doc):
    """Word 的中文字型要另外指定 eastAsia，不然會用預設的細明體。"""
    from docx.oxml.ns import qn
    for style_name in ("Normal", "List Bullet", "Intense Quote",
                       "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = CJK_FONT
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), CJK_FONT)


def markdown_to_docx_bytes(md: str, title: str = None) -> bytes:
    """回傳 .docx 的位元組內容。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _set_cjk_font(doc)
    if title:
        doc.core_properties.title = title

    md = _COMMENT_RE.sub("", md or "")
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue

        m = _H_RE.match(line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading("", level=level)
            par = doc.paragraphs[-1]
            _add_runs(par, m.group(2).strip())
            continue

        m = _QUOTE_RE.match(line)
        if m:
            par = doc.add_paragraph(style="Intense Quote")
            _add_runs(par, m.group(1).strip())
            continue

        m = _LI_RE.match(line)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            _add_runs(par, m.group(1).strip())
            continue

        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(6)
        _add_runs(par, line.strip())

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
